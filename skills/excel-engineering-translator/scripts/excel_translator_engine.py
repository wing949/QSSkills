import openpyxl
import re
import os
import json
from openpyxl.styles import Font, Alignment
from openpyxl.cell.cell import MergedCell

def has_cjk(text):
    if not isinstance(text, str):
        return False
    return bool(re.search(r'[\u4e00-\u9fff]', text))

def get_target_cell(ws, row, col):
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return cell
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col)
    return cell

def extract_unique_strings(excel_path, target_columns_per_sheet=None):
    wb = openpyxl.load_workbook(excel_path, data_only=False)
    unique_texts = set()
    indexed_tasks = []
    
    for s_idx, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        cols = target_columns_per_sheet.get(name, [(2, 3), (4, 5)]) if target_columns_per_sheet else [(2, 3), (4, 5)]
        for r in range(1, ws.max_row + 1):
            for c_src, c_tgt in cols:
                val = ws.cell(r, c_src).value
                if val and has_cjk(str(val)):
                    text = str(val).strip()
                    unique_texts.add(text)
                    indexed_tasks.append({'sheet': name, 'row': r, 'src_col': c_src, 'tgt_col': c_tgt, 'text': text})
                    
    return list(unique_texts), indexed_tasks

def apply_translations(excel_path, output_path, master_dict, indexed_tasks, font_name='Arial', font_size=12):
    wb = openpyxl.load_workbook(excel_path, data_only=False)
    font_apply = Font(name=font_name, size=font_size)
    align_apply = Alignment(wrap_text=True, vertical='center')
    
    updated_count = 0
    for t in indexed_tasks:
        ws = wb[t['sheet']]
        val_vn = master_dict.get(t['text'])
        if val_vn:
            cell = get_target_cell(ws, t['row'], t['tgt_col'])
            if not isinstance(cell, MergedCell):
                cell.value = val_vn
                cell.font = font_apply
                cell.alignment = align_apply
                updated_count += 1
                
    wb.save(output_path)
    print(f"Updated {updated_count} cells. Saved to {output_path}")

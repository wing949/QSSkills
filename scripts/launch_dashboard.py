"""
CSA AI OFFICE — Universal Project Dashboard Launcher & Data Bridge
Extracts live data from ANY 4-sheet CSA BOQ Excel Workbook (.xlsx)
and launches the interactive Virtual QS Office Dashboard.
"""

import sys
import os
import json
import webbrowser
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

DEFAULT_BOQ = r"C:\Users\TVCHUONG\Desktop\AI\11_BOQ\BOQ_B11_Chemical_Storage_OP2_Takeoff.xlsx"
HTML_TEMPLATE = r"C:\Users\TVCHUONG\Desktop\AI\11_BOQ\csa_ai_office_dashboard.html"


def parse_csa_workbook(file_path):
    """Trích xuất toàn diện dữ liệu từ bất kỳ file BOQ chuẩn CSA 4 sheet."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

    print(f"[1/3] Đang nạp và phân tích file Excel: {os.path.basename(file_path)}...")
    wb = openpyxl.load_workbook(file_path, data_only=True)

    data = {
        "file_name": os.path.basename(file_path),
        "file_path": file_path,
        "project_title": "DỰ ÁN XÂY DỰNG",
        "sub_title": "",
        "doc_info": "",
        "items": [],
        "sections": [],
        "qa_gates": [],
        "rfis": [],
        "earthwork": None,
        "summary": {
            "total_items": 0,
            "section_a_count": 0,
            "section_b_count": 0,
            "section_c_count": 0,
            "qa_pass_count": 0,
            "qa_total_count": 0,
            "rfi_count": 0,
            "rfi_resolved_count": 0,
        }
    }

    # 1. Parse Sheet BOQ
    if "BOQ" in wb.sheetnames:
        ws = wb["BOQ"]
        data["project_title"] = str(ws.cell(1, 1).value or "DỰ ÁN XÂY DỰNG")
        data["sub_title"] = str(ws.cell(2, 1).value or "")
        data["doc_info"] = str(ws.cell(3, 1).value or "")

        current_sec1 = ""
        current_sec2 = ""

        for r in range(6, ws.max_row + 1):
            val_a = ws.cell(r, 1).value
            val_b = ws.cell(r, 2).value
            val_c = ws.cell(r, 3).value
            val_d = ws.cell(r, 4).value
            val_e = ws.cell(r, 5).value

            if not val_a and not val_b:
                continue

            str_a = str(val_a).strip() if val_a else ""
            str_b = str(val_b).strip() if val_b else ""

            # Check Level 1 Section
            if str_a.startswith("PHẦN") or "SECTION" in str_a.upper() or (not val_c and not val_d and "." not in str_a):
                current_sec1 = str_b or str_a
                data["sections"].append({"code": str_a, "title": str_b, "level": 1})
                continue

            # Check Level 2 Subsection (e.g., A.I, B.II, C.I)
            if not val_c and not val_d and any(str_a.startswith(p) for p in ["A.", "B.", "C.", "D."]) and "." in str_a and len(str_a) <= 6:
                current_sec2 = str_b or str_a
                data["sections"].append({"code": str_a, "title": str_b, "level": 2})
                continue

            # Data Item
            if val_c and val_d is not None:
                item_obj = {
                    "wbs": str_a,
                    "desc": str_b,
                    "unit": str(val_c).strip(),
                    "qty": float(val_d) if isinstance(val_d, (int, float)) else 0.0,
                    "formula": str(val_e or "").strip(),
                    "sec1": current_sec1,
                    "sec2": current_sec2,
                }
                data["items"].append(item_obj)
                data["summary"]["total_items"] += 1

                if str_a.startswith("A."):
                    data["summary"]["section_a_count"] += 1
                elif str_a.startswith("B."):
                    data["summary"]["section_b_count"] += 1
                elif str_a.startswith("C."):
                    data["summary"]["section_c_count"] += 1

    # 2. Parse Sheet QA_Audit
    if "QA_Audit" in wb.sheetnames:
        ws_qa = wb["QA_Audit"]
        for r in range(6, ws_qa.max_row + 1):
            gate = ws_qa.cell(r, 1).value
            rule = ws_qa.cell(r, 2).value
            target = ws_qa.cell(r, 3).value
            sev = ws_qa.cell(r, 4).value
            evidence = ws_qa.cell(r, 5).value
            action = ws_qa.cell(r, 6).value
            owner = ws_qa.cell(r, 7).value
            status = ws_qa.cell(r, 8).value
            date = ws_qa.cell(r, 9).value

            if gate and rule and target:
                data["qa_gates"].append({
                    "gate": str(gate).strip(),
                    "rule": str(rule).strip(),
                    "target": str(target).strip(),
                    "severity": str(sev or "Major").strip(),
                    "evidence": str(evidence or "").strip(),
                    "action": str(action or "").strip(),
                    "owner": str(owner or "QA Lead").strip(),
                    "status": str(status or "Pass").strip(),
                    "date": str(date or "2026-09-24").strip()[:10],
                })
                data["summary"]["qa_total_count"] += 1
                if str(status).strip() == "Pass":
                    data["summary"]["qa_pass_count"] += 1

    # 3. Parse Sheet RFI
    if "RFI_Kien_Nghi_Bo_Sung" in wb.sheetnames:
        ws_rfi = wb["RFI_Kien_Nghi_Bo_Sung"]
        for r in range(6, ws_rfi.max_row + 1):
            rfi_id = ws_rfi.cell(r, 1).value
            trade = ws_rfi.cell(r, 2).value
            loc = ws_rfi.cell(r, 3).value
            qty = ws_rfi.cell(r, 4).value
            formula = ws_rfi.cell(r, 5).value
            content = ws_rfi.cell(r, 6).value
            prop = ws_rfi.cell(r, 7).value
            meta = ws_rfi.cell(r, 8).value

            if rfi_id and trade and content:
                data["rfis"].append({
                    "id": str(rfi_id).strip(),
                    "trade": str(trade).strip(),
                    "location": str(loc or "").strip(),
                    "qty": str(qty or "").strip(),
                    "formula": str(formula or "").strip(),
                    "content": str(content or "").strip(),
                    "proposal": str(prop or "").strip(),
                    "meta": str(meta or "").strip(),
                })
                data["summary"]["rfi_count"] += 1
                if "Status=Resolved" in str(meta) or "Status=Closed" in str(meta):
                    data["summary"]["rfi_resolved_count"] += 1

    # 4. Check Earthwork Balance
    if "Takeoff_Calculation" in wb.sheetnames:
        ws_calc = wb["Takeoff_Calculation"]
        if ws_calc.cell(6, 14).value:  # Column N
            data["earthwork"] = {
                "desc": str(ws_calc.cell(6, 14).value or ""),
                "ref": str(ws_calc.cell(6, 15).value or ""),
                "unit": str(ws_calc.cell(6, 17).value or "m3"),
                "cut": float(ws_calc.cell(6, 18).value or 0),
                "haul": float(ws_calc.cell(6, 19).value or 0),
                "fill": float(ws_calc.cell(6, 20).value or 0),
                "reuse": float(ws_calc.cell(6, 21).value or 0),
                "import": float(ws_calc.cell(6, 22).value or 0),
                "surplus": float(ws_calc.cell(6, 23).value or 0),
            }

    print(f"[2/3] Trích xuất thành công: {data['summary']['total_items']} items BOQ, {data['summary']['qa_total_count']} Gates QA/QC, {data['summary']['rfi_count']} RFIs.")
    return data


def inject_and_launch_dashboard(data, html_template_path):
    """Nhúng dữ liệu động vào file Dashboard và mở trình duyệt."""
    print(f"[3/3] Đang cập nhật dữ liệu động vào Dashboard...")
    if not os.path.exists(html_template_path):
        raise FileNotFoundError(f"Không tìm thấy template: {html_template_path}")

    with open(html_template_path, "r", encoding="utf-8") as f:
        html_content = f.read()

    # Nhúng biến toàn cục window.DYNAMIC_PROJECT_DATA
    json_str = json.dumps(data, ensure_ascii=False, indent=2)
    injection_script = f"\n<script id=\"injected-project-data\">\n  window.DYNAMIC_PROJECT_DATA = {json_str};\n</script>\n"

    # Thay thế hoặc chèn trước </head>
    if '<script id="injected-project-data">' in html_content:
        import re
        html_content = re.sub(
            r'<script id="injected-project-data">.*?</script>',
            injection_script.strip(),
            html_content,
            flags=re.DOTALL
        )
    else:
        html_content = html_content.replace("</head>", f"{injection_script}</head>")

    with open(html_template_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    print(f"Cập nhật thành công! Đang khởi chạy Dashboard trên trình duyệt mặc định...")
    webbrowser.open(f"file:///{os.path.abspath(html_template_path)}")
    print(f"URL: file:///{os.path.abspath(html_template_path)}")


if __name__ == "__main__":
    target_file = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_BOQ
    try:
        project_data = parse_csa_workbook(target_file)
        inject_and_launch_dashboard(project_data, HTML_TEMPLATE)
        print("\nHOÀN TẤT: Dashboard đã kết nối và hiển thị dữ liệu của dự án mới!")
    except Exception as e:
        print(f"\n[LỖI]: {e}", file=sys.stderr)
        sys.exit(1)
